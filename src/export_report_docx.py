"""Phase 7 - convert docs/PROJECT_REPORT.md to a submission-ready .docx.

Deliberately a small, dependency-light converter rather than a general Markdown
engine. It handles exactly the constructs the report uses:

  # / ## / ### headings, paragraphs, **bold**, `code`, *italic*,
  bullet lists, numbered lists, pipe tables, fenced code blocks,
  --- horizontal rules, and ![alt](path) images.

Run:  python src/export_report_docx.py
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "docs" / "PROJECT_REPORT.md"
OUT = ROOT / "docs" / "PROJECT_REPORT.docx"

INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*]+\*|\[.+?\]\(.+?\))")
FENCE = re.compile(r"^```(\w*)\s*$")


def add_runs(paragraph, text: str) -> None:
    """Render inline Markdown emphasis into runs."""
    pos = 0
    for match in INLINE.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos:match.start()])
        token = match.group(0)
        if token.startswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0xB0, 0x30, 0x60)
        elif token.startswith("["):
            label, url = re.match(r"\[(.+?)\]\((.+?)\)", token).groups()  # type: ignore[union-attr]
            run = paragraph.add_run(f"{label} ({url})" if label != url else url)
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x1A, 0x4F, 0xA0)
            run.underline = True
        else:
            paragraph.add_run(token[1:-1]).italic = True
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def split_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def main() -> int:
    if not SRC.exists():
        raise SystemExit(f"missing {SRC}")

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Inches(0.8)
        sec.top_margin = sec.bottom_margin = Inches(0.8)

    lines = SRC.read_text(encoding="utf-8").splitlines()
    i = 0
    in_code = False
    code_lang = ""
    code_buf: list[str] = []
    first_heading_done = False

    while i < len(lines):
        raw = lines[i]
        stripped = raw.strip()

        # ---- fenced code -------------------------------------------------
        fence = FENCE.match(stripped)
        if fence and not in_code:
            in_code, code_lang = True, fence.group(1)
            code_buf = []
            i += 1
            continue
        if in_code:
            if stripped == "```":
                para = doc.add_paragraph()
                run = para.add_run("\n".join(code_buf))
                run.font.name = "Consolas"
                run.font.size = Pt(8.5)
                para.paragraph_format.left_indent = Inches(0.25)
                para.paragraph_format.space_after = Pt(8)
                in_code = False
            else:
                code_buf.append(raw)
            i += 1
            continue

        # ---- image -------------------------------------------------------
        img = re.match(r"!\[(.*?)\]\((.+?)\)", stripped)
        if img:
            target = img.group(2)
            # Report-relative first (docs/architecture.png), then repo root.
            for candidate in (SRC.parent / target, ROOT / target):
                if candidate.exists():
                    doc.add_picture(str(candidate), width=Inches(6.8))
                    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cap = doc.add_paragraph(img.group(1))
                    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in cap.runs:
                        r.italic = True
                        r.font.size = Pt(9)
                        r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
                    break
            else:
                doc.add_paragraph(f"[missing image: {target}]").italic = True
            i += 1
            continue

        # ---- table -------------------------------------------------------
        if stripped.startswith("|") and i + 1 < len(lines) and re.match(
            r"^\|[\s:|-]+\|$", lines[i + 1].strip()
        ):
            header = split_row(stripped)
            i += 2
            body: list[list[str]] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                body.append(split_row(lines[i].strip()))
                i += 1
            table = doc.add_table(rows=1, cols=len(header))
            table.style = "Light Grid Accent 1"
            for cell, text in zip(table.rows[0].cells, header):
                cell.text = ""
                run = cell.paragraphs[0].add_run(text)
                run.bold = True
                run.font.size = Pt(9.5)
            for row in body:
                cells = table.add_row().cells
                for cell, text in zip(cells, row):
                    cell.text = ""
                    add_runs(cell.paragraphs[0], text)
                    for r in cell.paragraphs[0].runs:
                        if r.font.size is None:
                            r.font.size = Pt(9.5)
            doc.add_paragraph()
            continue

        # ---- headings ----------------------------------------------------
        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            text = stripped[level:].strip()
            if level == 1 and not first_heading_done:
                head = doc.add_heading(text, level=0)
                head.alignment = WD_ALIGN_PARAGRAPH.CENTER
                first_heading_done = True
            else:
                doc.add_heading(text, level=min(level, 4))
            i += 1
            continue

        # ---- rules -------------------------------------------------------
        if stripped in {"---", "***", "___"}:
            i += 1
            continue

        # ---- lists -------------------------------------------------------
        bullet = re.match(r"^[-*]\s+(.*)$", stripped)
        numbered = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if bullet:
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, bullet.group(1))
            i += 1
            continue
        if numbered:
            p = doc.add_paragraph(style="List Number")
            add_runs(p, numbered.group(2))
            i += 1
            continue

        # ---- blockquote --------------------------------------------------
        if stripped.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            add_runs(p, stripped.lstrip("> ").strip())
            for r in p.runs:
                r.italic = True
            i += 1
            continue

        # ---- blank -------------------------------------------------------
        if not stripped:
            i += 1
            continue

        # ---- paragraph ---------------------------------------------------
        p = doc.add_paragraph()
        add_runs(p, stripped)
        i += 1

    doc.save(OUT)
    size = OUT.stat().st_size
    print(f"wrote {OUT.relative_to(ROOT)}  ({size} bytes)")
    if size < 20000:
        print("WARNING: output looks too small - check the source parsed correctly")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
